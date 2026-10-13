# Módulo de RAG - VERSÃO OTIMIZADA COM OCR E PARALELIZAÇÃO

import sys
import os
import json
import hashlib
import multiprocessing
from concurrent.futures import ProcessPoolExecutor, as_completed

import docx
import PyPDF2
from pptx import Presentation
from os import listdir
from os.path import isfile, join, isdir

from langchain_text_splitters import TokenTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams
from langchain_qdrant import Qdrant

# OCR
import pytesseract
from pdf2image import convert_from_path
from PIL import Image

# Barra de progresso
from tqdm import tqdm

# Filtra warnings
import warnings
warnings.filterwarnings('ignore')


# ============================================================
# CONFIGURAÇÕES GLOBAIS
# ============================================================

# Modelo de embedding
MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
VECTOR_SIZE = 384
# Alternativa mais precisa (mas 5x mais lenta):
# MODEL_NAME = "sentence-transformers/msmarco-bert-base-dot-v5"
# VECTOR_SIZE = 768

# Qdrant
QDRANT_URL = "http://localhost:6333"
COLLECTION_NAME = "ragVectorDB"

# Text splitter
CHUNK_SIZE = 800
CHUNK_OVERLAP = 80

# OCR
OCR_IDIOMA = "por+eng"   # use apenas "por" se todos os docs são em português (mais rápido)
OCR_DPI = 200            # aumente para 300 se qualidade ruim

# Paralelização — use nproc-1 para não travar o sistema
NUM_WORKERS = max(1, multiprocessing.cpu_count() - 1)

# Controle de indexação incremental
CONTROLE_FILE = "arquivos_indexados.json"

# Extensões suportadas
EXTENSOES_VALIDAS = (
    ".pdf", ".txt", ".docx", ".pptx",
    ".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"
)


# ============================================================
# CONTROLE DE INDEXAÇÃO INCREMENTAL
# ============================================================

def carregar_controle():
    """Carrega lista de arquivos já indexados (com hash para detectar mudanças)."""
    if os.path.exists(CONTROLE_FILE):
        with open(CONTROLE_FILE, "r") as f:
            return json.load(f)
    return {}


def salvar_controle(controle):
    with open(CONTROLE_FILE, "w") as f:
        json.dump(controle, f, indent=2)


def hash_arquivo(caminho):
    """Gera hash MD5 do arquivo para detectar se mudou."""
    h = hashlib.md5()
    with open(caminho, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


# ============================================================
# LISTAGEM DE ARQUIVOS
# ============================================================

def rag_lista_arquivos(dir):
    """Lista todos os arquivos em um diretório (recursivo)."""
    arquivo_list = []
    for f in listdir(dir):
        if isfile(join(dir, f)):
            arquivo_list.append(join(dir, f))
        elif isdir(join(dir, f)):
            arquivo_list += rag_lista_arquivos(join(dir, f))
    return arquivo_list


# ============================================================
# EXTRATORES DE TEXTO POR FORMATO
# ============================================================

def rag_carrega_texto_word(arquivoname):
    """Extrai texto de arquivo .docx"""
    doc = docx.Document(arquivoname)
    fullText = [para.text for para in doc.paragraphs]
    return '\n'.join(fullText)


def rag_carrega_texto_pptx(arquivoname):
    """Extrai texto de arquivo .pptx"""
    prs = Presentation(arquivoname)
    fullText = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                fullText.append(shape.text)
    return '\n'.join(fullText)


# ============================================================
# OCR — para PDFs escaneados e imagens
# ============================================================

def rag_ocr_imagem(caminho, idioma=OCR_IDIOMA):
    """Faz OCR de um arquivo de imagem (jpg, png, tif, bmp)."""
    try:
        img = Image.open(caminho)
        return pytesseract.image_to_string(img, lang=idioma)
    except Exception as e:
        print(f"  [OCR-ERRO] {caminho}: {e}")
        return ""


def rag_ocr_pdf(caminho, idioma=OCR_IDIOMA, dpi=OCR_DPI):
    """Converte PDF em imagens e faz OCR em cada página."""
    try:
        paginas = convert_from_path(caminho, dpi=dpi)
        textos = []
        for pagina in paginas:
            texto = pytesseract.image_to_string(pagina, lang=idioma)
            textos.append(texto)
        return "\n".join(textos)
    except Exception as e:
        print(f"  [OCR-PDF-ERRO] {caminho}: {e}")
        return ""


# ============================================================
# EXTRAÇÃO UNIFICADA (com fallback para OCR)
# ============================================================

def extrair_conteudo(arquivo):
    """
    Extrai conteúdo de um arquivo. Se for PDF escaneado ou imagem,
    aplica OCR automaticamente.
    """
    nome_lower = arquivo.lower()

    # PDF: tenta extração normal; se vier vazio, faz OCR
    if nome_lower.endswith(".pdf"):
        try:
            reader = PyPDF2.PdfReader(arquivo)
            texto = " ".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            texto = ""

        # Se veio pouco texto, provavelmente é PDF escaneado → OCR
        if len(texto.strip()) < 50:
            texto = rag_ocr_pdf(arquivo)
        return texto

    elif nome_lower.endswith(".txt"):
        with open(arquivo, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    elif nome_lower.endswith(".docx"):
        return rag_carrega_texto_word(arquivo)

    elif nome_lower.endswith(".pptx"):
        return rag_carrega_texto_pptx(arquivo)

    elif nome_lower.endswith((".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp")):
        return rag_ocr_imagem(arquivo)

    return None


# ============================================================
# WORKER PARA PROCESSAMENTO PARALELO
# ============================================================

def processar_arquivo_worker(args):
    """
    Worker que roda em processo separado.
    Faz extração + split e retorna os chunks prontos para envio ao Qdrant.
    
    Retorna: (arquivo, hash, lista_de_chunks, status)
      status pode ser: "OK", "VAZIO", "ERRO: <msg>"
    """
    arquivo, h = args
    try:
        conteudo = extrair_conteudo(arquivo)
        
        if not conteudo or not conteudo.strip():
            return (arquivo, h, [], "VAZIO")
        
        # Splitter precisa ser criado dentro do worker
        splitter = TokenTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
        textos = splitter.split_text(conteudo)
        
        return (arquivo, h, textos, "OK")
    
    except Exception as e:
        return (arquivo, h, [], f"ERRO: {e}")


# ============================================================
# INDEXAÇÃO PRINCIPAL
# ============================================================

def main_indexing(mypath):

    # -----------------------------------------------------
    # Modelo de embedding
    # -----------------------------------------------------
    print(f"Carregando modelo de embedding: {MODEL_NAME}")
    
    hf = HuggingFaceEmbeddings(
        model_name=MODEL_NAME,
        model_kwargs={'device': 'cpu'},
        encode_kwargs={
            'normalize_embeddings': True,
            'batch_size': 64  # processa 64 chunks por vez
        }
    )

    # -----------------------------------------------------
    # Cliente Qdrant
    # -----------------------------------------------------
    client = QdrantClient(QDRANT_URL)

    if not client.collection_exists(COLLECTION_NAME):
        client.create_collection(
            COLLECTION_NAME,
            vectors_config=VectorParams(size=VECTOR_SIZE, distance=Distance.COSINE)
        )
        print(f"Coleção '{COLLECTION_NAME}' criada com {VECTOR_SIZE} dimensões.")
    else:
        print(f"Coleção '{COLLECTION_NAME}' já existe. Indexação incremental ativa.")

    qdrant = Qdrant(client, COLLECTION_NAME, hf)

    # -----------------------------------------------------
    # Carrega controle e lista arquivos
    # -----------------------------------------------------
    controle = carregar_controle()
    lista_arquivos = rag_lista_arquivos(mypath)
    print(f"\nTotal de arquivos encontrados: {len(lista_arquivos)}")

    # Filtra apenas arquivos NOVOS ou MODIFICADOS
    arquivos_para_processar = []
    for arquivo in lista_arquivos:
        if not arquivo.lower().endswith(EXTENSOES_VALIDAS):
            continue
        h = hash_arquivo(arquivo)
        if controle.get(arquivo) == h:
            continue  # já indexado e sem mudanças
        arquivos_para_processar.append((arquivo, h))

    print(f"Arquivos novos/modificados para indexar: {len(arquivos_para_processar)}")

    if not arquivos_para_processar:
        print("Nada para fazer. Tudo já indexado! ✅")
        return

    print(f"Usando {NUM_WORKERS} processos paralelos.\n")

    # -----------------------------------------------------
    # Loop principal PARALELIZADO
    # -----------------------------------------------------
    # Contadores de status
    contador = {"ok": 0, "vazio": 0, "erro": 0, "erro_qdrant": 0}

    with ProcessPoolExecutor(max_workers=NUM_WORKERS) as executor:
        # Submete todos os arquivos para processamento paralelo
        futures = {
            executor.submit(processar_arquivo_worker, args): args
            for args in arquivos_para_processar
        }

        # Barra de progresso conforme workers vão terminando
        pbar = tqdm(as_completed(futures), total=len(futures), desc="Arquivos")
        
        for future in pbar:
            try:
                arquivo, h, textos, status = future.result()
            except Exception as e:
                print(f"  [ERRO-WORKER] {e}")
                contador["erro"] += 1
                continue

            # Nome curto para exibir na barra
            nome_curto = os.path.basename(arquivo)[:40]
            pbar.set_postfix_str(f"último: {nome_curto}")

            if status == "VAZIO":
                # Marca como processado pra não tentar de novo
                controle[arquivo] = h
                salvar_controle(controle)
                contador["vazio"] += 1
                continue

            if status.startswith("ERRO"):
                print(f"\n  [{status}] {arquivo}")
                contador["erro"] += 1
                continue

            # ---- Envio ao Qdrant (sequencial, no processo principal) ----
            try:
                metadata = [{"path": arquivo} for _ in textos]
                qdrant.add_texts(textos, metadatas=metadata)
                controle[arquivo] = h
                salvar_controle(controle)
                contador["ok"] += 1
            except Exception as e:
                print(f"\n  [ERRO-QDRANT] {arquivo}: {e}")
                contador["erro_qdrant"] += 1

    # -----------------------------------------------------
    # Relatório final
    # -----------------------------------------------------
    print("\n" + "=" * 50)
    print("✅ Indexação Concluída!")
    print("=" * 50)
    print(f"  Sucesso:            {contador['ok']}")
    print(f"  Vazios (sem texto): {contador['vazio']}")
    print(f"  Erros extração:     {contador['erro']}")
    print(f"  Erros Qdrant:       {contador['erro_qdrant']}")
    print(f"  Total processado:   {sum(contador.values())}")
    print("=" * 50 + "\n")


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    arguments = sys.argv
    if len(arguments) > 1:
        main_indexing(arguments[1])
    else:
        print("Uso: python rag_rag.py <caminho_da_pasta>")
        print("Exemplo: python rag_rag.py Documentos")